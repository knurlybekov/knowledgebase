"""Fixtures for pipeline tests: source documents, events, a static fetcher,
small Office and PDF files built in memory."""

from __future__ import annotations

import io
from datetime import UTC, datetime

from knowledge.api import ChangeEvent, DownloadResult, SourceDocument
from organizations.api import Principal


def make_source(
    source_id: str,
    label: str,
    title: str = "Doc",
    mime: str = "text/markdown",
    source_type: str = "google_drive",
) -> SourceDocument:
    return SourceDocument(
        source_type=source_type,
        source_id=source_id,
        version_label=label,
        title=title,
        mime_type=mime,
        source_url=f"https://docs.google.com/document/d/{source_id}",
        modified_at=datetime(2026, 1, 1, tzinfo=UTC),
    )


def make_event(
    seq: int,
    sd: SourceDocument | None,
    kind: str = "upsert",
    source_id: str | None = None,
    generation: int | None = None,
    observed_at: datetime | None = None,
) -> ChangeEvent:
    return ChangeEvent(
        kind=kind,
        source_id=source_id or (sd.source_id if sd else ""),
        seq=seq,
        observed_at=observed_at or datetime.now(UTC),
        reconcile_generation=generation,
        document=sd,
    )


class StaticFetcher:
    """Returns bytes by version label and counts calls. `on_fetch` runs inside
    the staging phase to simulate concurrent events."""

    def __init__(self, contents: dict[str, bytes], mime: str = "text/markdown", on_fetch=None):
        self.contents = contents
        self.mime = mime
        self.calls = 0
        self.on_fetch = on_fetch

    def __call__(self, sd: SourceDocument) -> DownloadResult:
        self.calls += 1
        if self.on_fetch:
            self.on_fetch(sd)
        return DownloadResult(
            content=self.contents[sd.version_label],
            version_label=sd.version_label,
            mime_type=self.mime,
        )


def principal_for(org, user, role: str = "member") -> Principal:
    return Principal(user=user, organization=org, role=role)


def docx_bytes(paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """paragraphs: (style, text) with style 'Heading 1', 'Heading 2' or 'Normal'."""
    import docx

    d = docx.Document()
    for style, text in paragraphs:
        d.add_paragraph(text, style=style if style != "Normal" else None)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: list[tuple[str, str]]) -> bytes:
    from pptx import Presentation

    p = Presentation()
    layout = p.slide_layouts[1]
    for title, body in slides:
        s = p.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(pages: list[str]) -> bytes:
    """A minimal valid PDF with one Helvetica text line per page."""
    objects: list[bytes] = []
    kids = []
    page_numbers = []
    n_fixed = 3  # catalog, pages, font
    for i, text in enumerate(pages):
        page_obj = n_fixed + 1 + i * 2
        content_obj = page_obj + 1
        kids.append(f"{page_obj} 0 R")
        page_numbers.append((page_obj, content_obj, text))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for _page_obj, content_obj, text in page_numbers:
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_obj} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode()
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def count_visible(principal) -> int:
    """visible_documents() inside the organization scope, as the web layer calls it."""
    from infra.api import with_org
    from knowledge.api import visible_documents

    with with_org(principal.organization_id):
        return visible_documents(principal).count()


def visible_labels(principal) -> list[str]:
    from infra.api import with_org
    from knowledge.api import visible_versions

    with with_org(principal.organization_id):
        return sorted(visible_versions(principal).values_list("version_label", flat=True))
