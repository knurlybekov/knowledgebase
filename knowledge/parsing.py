"""Document parsers. Each returns sections with heading paths and locators so
a chunk traces back to a page, slide, sheet or heading."""

from __future__ import annotations

import html
import io
import re

from .types import ParsedDocument, Section, UnsupportedFormat

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
TEXT_TYPES = {"text/plain", "text/markdown", "text/csv", "text/x-markdown"}
MAX_SHEET_ROWS = 2000
MAX_TEXT_BYTES = 20 * 1024 * 1024

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def is_supported(mime_type: str, filename: str = "") -> bool:
    """True when `parse` has a parser for this type. Checked before a download so
    audio, video, archives and binaries are quarantined without a transfer."""
    mime = (mime_type or "").split(";")[0].strip().lower()
    name = (filename or "").lower()
    if mime in (DOCX, PPTX, XLSX, PDF) or mime in TEXT_TYPES or mime.startswith("text/"):
        return True
    return name.endswith((".docx", ".pptx", ".xlsx", ".pdf", ".md", ".txt", ".markdown"))


def parse(content: bytes, mime_type: str, filename: str = "") -> ParsedDocument:
    """Raises UnsupportedFormat for an unknown type and for content the parser
    for that type cannot read (truncated zip, empty PDF, placeholder files):
    both are terminal for the revision, never a retry."""
    mime = (mime_type or "").split(";")[0].strip().lower()
    name = (filename or "").lower()
    try:
        return _dispatch(content, mime, name)
    except UnsupportedFormat:
        raise
    except Exception as exc:
        raise UnsupportedFormat(
            f"unreadable {mime or name or 'unknown'}: {type(exc).__name__}"
        ) from exc


def _dispatch(content: bytes, mime: str, name: str) -> ParsedDocument:
    if mime == DOCX or name.endswith(".docx"):
        return _parse_docx(content)
    if mime == PPTX or name.endswith(".pptx"):
        return _parse_pptx(content)
    if mime == XLSX or name.endswith(".xlsx"):
        return _parse_xlsx(content)
    if mime == PDF or name.endswith(".pdf"):
        return _parse_pdf(content)
    if (
        mime in TEXT_TYPES
        or mime.startswith("text/")
        or name.endswith((".md", ".txt", ".markdown"))
    ):
        return _parse_text(content)
    raise UnsupportedFormat(mime or name or "unknown")


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _parse_text(content: bytes) -> ParsedDocument:
    if len(content) > MAX_TEXT_BYTES:
        raise UnsupportedFormat("text file too large")
    text = content.decode("utf-8", errors="replace")
    sections: list[Section] = []
    heading_path: list[str] = []
    buffer: list[str] = []

    def flush():
        body = _clean("\n".join(buffer))
        if body:
            sections.append(Section(tuple(heading_path), body, {}))
        buffer.clear()

    for line in text.splitlines():
        m = _MD_HEADING.match(line)
        if m:
            flush()
            level = len(m.group(1))
            heading_path = [*heading_path[: level - 1], m.group(2).strip()]
            continue
        buffer.append(line)
    flush()
    if not sections and text.strip():
        sections.append(Section((), _clean(text), {}))
    return ParsedDocument(sections, "text")


def _parse_docx(content: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        # Live drive 2026-10-16: a 32 MB plan with a broken internal relationship
        # (KeyError on a missing part). Fall back to the raw document XML.
        return _parse_docx_xml(content, reason=type(exc).__name__)
    sections: list[Section] = []
    heading_path: list[str] = []
    buffer: list[str] = []

    def flush():
        body = _clean("\n".join(buffer))
        if body:
            sections.append(Section(tuple(heading_path), body, {}))
        buffer.clear()

    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if style.lower().startswith("heading") and text:
            flush()
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            heading_path = [*heading_path[: level - 1], text]
            continue
        if text:
            buffer.append(text)
    flush()
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append(Section(tuple(heading_path), "\n".join(rows), {"table": index}))
    return ParsedDocument(sections, "docx")


def _parse_docx_xml(content: bytes, *, reason: str) -> ParsedDocument:
    """Paragraph text from word/document.xml without python-docx: one section,
    no heading structure. Raises UnsupportedFormat when even that is unreadable."""
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", "replace")
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise UnsupportedFormat(f"docx: {reason}") from exc
    paragraphs = []
    for para in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S):
        text = "".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", para))
        text = html.unescape(text).strip()
        if text:
            paragraphs.append(text)
    body = _clean("\n".join(paragraphs))
    if not body:
        raise UnsupportedFormat(f"docx: {reason}")
    return ParsedDocument([Section((), body, {"fallback": "document.xml"})], "docx-xml")


def _parse_pptx(content: bytes) -> ParsedDocument:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    sections: list[Section] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        parts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape == getattr(slide.shapes, "title", None):
                title = text
            else:
                parts.append(text)
        body = _clean("\n".join(parts))
        if title or body:
            sections.append(Section((title,) if title else (), body or title, {"slide": number}))
    return ParsedDocument(sections, "pptx")


def _parse_xlsx(content: bytes) -> ParsedDocument:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sections: list[Section] = []
    for sheet in workbook.worksheets:
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        if rows:
            sections.append(Section((sheet.title,), "\n".join(rows), {"sheet": sheet.title}))
    return ParsedDocument(sections, "xlsx")


def _parse_pdf(content: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    sections: list[Section] = []
    for number, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            sections.append(Section((), text, {"page": number}))
    return ParsedDocument(sections, "pdf")
