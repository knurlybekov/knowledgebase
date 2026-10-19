# ruff: noqa: E501  data file: long synthetic sentences
"""Synthetic study corpus for the academic demo: lecture notes, a slide deck and a
report, generated from text at seed time. No real course material."""

from __future__ import annotations

import io
from pathlib import Path

ROOT = Path(__file__).parent

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slides = [
        ("Operating Systems, Lecture 9: Scheduling", "Goals: throughput, turnaround, response time, fairness."),
        ("First-come first-served", "Simple and fair in arrival order. Suffers the convoy effect: one long CPU burst delays every short job behind it."),
        ("Shortest job first", "Optimal mean waiting time when burst lengths are known. Burst lengths are predicted with an exponential average; the course uses alpha = 0.5."),
        ("Round robin", "Time quantum q. A quantum of 10 to 100 ms keeps response time low; a quantum below the context-switch cost wastes the CPU."),
        ("Multilevel feedback queue", "New jobs start in the highest-priority queue; a job that uses its full quantum moves down. This is the scheduler the lab simulates in Assignment 4."),
    ]
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.placeholders[1].width = Inches(8)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _report() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Statistics project report: survey of commute times", 1)
    d.add_heading("Method", 2)
    d.add_paragraph(
        "We collected 212 responses from students in two faculties. Commute time was recorded in minutes. "
        "We compared the two faculties with Welch's t-test because the variances differed (Levene's test p = 0.01)."
    )
    d.add_heading("Result", 2)
    d.add_paragraph(
        "Mean commute time was 34.2 minutes in Science and 41.7 minutes in Business. The difference was significant "
        "(t = 2.91, p = 0.004, 95 percent confidence interval 2.4 to 12.6 minutes)."
    )
    d.add_heading("Limitation", 2)
    d.add_paragraph(
        "Responses were self-reported and the sample was not random. The instructor's feedback on the draft asked for an effect size, "
        "which we added as Cohen's d = 0.41."
    )
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def documents() -> dict[str, tuple[bytes, str, str]]:
    """{source_id: (bytes, mime_type, title)} for every synthetic study document."""
    out = {}
    for path in sorted(ROOT.glob("*.md")):
        out[f"gdrive:{path.stem}"] = (path.read_bytes(), "text/markdown", path.stem.replace("-", " ").title())
    out["gdrive:operating-systems-lecture-09-scheduling"] = (_deck(), PPTX, "Operating Systems Lecture 09 Scheduling.pptx")
    out["gdrive:statistics-project-report"] = (_report(), DOCX, "Statistics Project Report.docx")
    return out
